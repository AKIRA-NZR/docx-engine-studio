"""
Testes Automatizados para a arquitetura docx_engine:
- Validação e detecção de pendências
- Preenchimento e clonagem de propriedades XML (bordas e margens)
- Conversão DOCX -> PDF e PDF -> DOCX
- Resiliência contra bloqueio de arquivo (PermissionError)
"""

import os
import tempfile
from pathlib import Path
import pytest
import docx

from docx_engine import (
    DocumentEditor,
    DocumentValidator,
    clone_cell_properties,
    delete_paragraph_element,
    docx_to_pdf,
    pdf_to_docx
)
from configs.exemplo_relatorio import RECIPE_EXEMPLO

TEMPLATE_FILE = "Modelo_Relatorio_de_Evolucao_de_Front-End.docx"


def test_template_loading_and_inspection():
    """Garante que o DocumentEditor carrega o modelo e expõe parágrafos e tabelas."""
    editor = DocumentEditor(TEMPLATE_FILE)
    assert len(editor.doc.paragraphs) > 0
    assert len(editor.doc.tables) == 11


def test_validator_detects_unfilled_template():
    """Verifica se o validador acusa corretamente pendências no modelo vazio."""
    editor = DocumentEditor(TEMPLATE_FILE)
    report = editor.validate()
    # No modelo não preenchido, deve haver erros de placeholders (____)
    assert not report.is_valid
    assert len(report.errors) > 0
    assert any("____" in err.snippet for err in report.errors)


def test_filling_and_validation_approval():
    """Aplica a receita e verifica se o documento preenchido passa na auditoria."""
    with tempfile.TemporaryDirectory() as tmpdir:
        out_docx = Path(tmpdir) / "teste_preenchido.docx"
        editor = DocumentEditor(TEMPLATE_FILE)
        editor.apply_recipe(RECIPE_EXEMPLO)
        saved = editor.save(str(out_docx))
        assert Path(saved).exists()

        # Auditoria do documento preenchido
        report = editor.validate()
        assert report.is_valid
        assert len(report.errors) == 0


def test_grid_table_border_cloning():
    """Valida se as novas linhas adicionadas em tabelas matriciais mantêm 100% das bordas."""
    from docx.oxml.ns import qn
    with tempfile.TemporaryDirectory() as tmpdir:
        out_docx = Path(tmpdir) / "teste_bordas.docx"
        editor = DocumentEditor(TEMPLATE_FILE)
        editor.apply_recipe(RECIPE_EXEMPLO)
        editor.save(str(out_docx))

        doc = docx.Document(str(out_docx))
        # Verifica a Tabela 6 (Endpoints - que adiciona 8 novas linhas)
        table6 = doc.tables[6]
        assert len(table6.rows) == 15
        for r_idx, row in enumerate(table6.rows):
            for c_idx, cell in enumerate(row.cells):
                tcBorders = cell._tc.get_or_add_tcPr().find(qn('w:tcBorders'))
                assert tcBorders is not None, f"Borda ausente na Tabela 6, Linha {r_idx}, Coluna {c_idx}"


def test_safe_save_fallback_on_locked_file():
    """Testa se o mecanismo de salvamento seguro grava com sufixo _novo em caso de bloqueio (PermissionError)."""
    import stat
    with tempfile.TemporaryDirectory() as tmpdir:
        target_file = Path(tmpdir) / "documento_travado.docx"
        target_file.touch()

        # Torna o arquivo somente leitura para forçar PermissionError no Windows
        os.chmod(str(target_file), stat.S_IREAD)
        try:
            editor = DocumentEditor(TEMPLATE_FILE)
            saved_path = editor.save(str(target_file))
            # Deve salvar no arquivo alternativo com sufixo _novo.docx
            assert saved_path != str(target_file)
            assert "_novo" in saved_path
            assert Path(saved_path).exists()
        finally:
            # Restaura permissões para limpeza do diretório temporário
            os.chmod(str(target_file), stat.S_IWRITE)


def test_pdf_conversion_cycle():
    """Testa a exportação para PDF e a reconversão do PDF para DOCX."""
    with tempfile.TemporaryDirectory() as tmpdir:
        # 1. Cria um DOCX simples para teste
        doc = docx.Document()
        doc.add_heading("Teste de Conversão", 0)
        doc.add_paragraph("Este é um parágrafo de teste automatizado.")
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Item"
        table.rows[0].cells[1].text = "Status"
        table.rows[1].cells[0].text = "Motor"
        table.rows[1].cells[1].text = "Ativo"

        test_docx = Path(tmpdir) / "origem.docx"
        doc.save(str(test_docx))

        # 2. Converte DOCX -> PDF
        out_pdf = Path(tmpdir) / "saida.pdf"
        try:
            pdf_path = docx_to_pdf(str(test_docx), str(out_pdf))
            assert Path(pdf_path).exists()
            assert Path(pdf_path).stat().st_size > 0

            # 3. Converte PDF -> DOCX (pdf2docx)
            reconverted_docx = Path(tmpdir) / "reconvertido.docx"
            res_docx = pdf_to_docx(pdf_path, str(reconverted_docx))
            assert Path(res_docx).exists()
            assert Path(res_docx).stat().st_size > 0

            # 4. Verifica se o DOCX reconvertido abre normalmente no python-docx
            reopened = docx.Document(res_docx)
            assert len(reopened.paragraphs) > 0
        except RuntimeError as e:
            pytest.skip(f"Conversão para PDF indisponível no ambiente de teste: {e}")
