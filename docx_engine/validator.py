"""
Módulo de Validação e Auditoria de Documentos DOCX.
Verifica se restam campos não preenchidos, linhas pontilhadas residuais (____),
tags de template pendentes ({{...}}) e células de tabela esquecidas.
"""

import re
from dataclasses import dataclass, field
from typing import List, Dict, Any, Optional
import docx


@dataclass
class ValidationIssue:
    category: str       # 'placeholder', 'tag', 'empty_cell', 'suspicious'
    location: str       # Ex: 'Parágrafo 12' ou 'Tabela 3 (Linha 2, Coluna 1)'
    message: str
    severity: str = "warning"  # 'warning' ou 'error'
    snippet: str = ""


@dataclass
class ValidationReport:
    total_issues: int = 0
    errors: List[ValidationIssue] = field(default_factory=list)
    warnings: List[ValidationIssue] = field(default_factory=list)

    @property
    def is_valid(self) -> bool:
        return len(self.errors) == 0

    def add_issue(self, issue: ValidationIssue):
        self.total_issues += 1
        if issue.severity == "error":
            self.errors.append(issue)
        else:
            self.warnings.append(issue)

    def summary(self) -> str:
        status = "APROVADO" if self.is_valid else "PENDENTE DE PREENCHIMENTO"
        lines = [
            f"=== RELATÓRIO DE AUDITORIA DO DOCUMENTO: {status} ===",
            f"Total de pendências encontradas: {self.total_issues}",
            f"Erros críticos: {len(self.errors)}",
            f"Avisos / Alertas: {len(self.warnings)}",
            ""
        ]

        if self.errors:
            lines.append("--- ERROS CRÍTICOS (Preenchimento incompleto) ---")
            for err in self.errors:
                lines.append(f"  • [{err.location}] {err.message} -> {repr(err.snippet)}")
            lines.append("")

        if self.warnings:
            lines.append("--- AVISOS (Possíveis pendências) ---")
            for warn in self.warnings:
                lines.append(f"  • [{warn.location}] {warn.message} -> {repr(warn.snippet)}")
            lines.append("")

        return "\n".join(lines)


class DocumentValidator:
    """
    Auditor inteligente de conformidade para documentos DOCX.
    """

    def __init__(self, doc: docx.Document):
        self.doc = doc

    def validate(
        self,
        check_underscores: bool = True,
        check_tags: bool = True,
        check_empty_cells: bool = False,
        ignore_table_indices: Optional[List[int]] = None,
        ignore_signatures: bool = True
    ) -> ValidationReport:
        report = ValidationReport()
        ignore_tables = ignore_table_indices or [0]  # Tabela 0 geralmente é instrução

        # 1. Auditoria de Parágrafos
        # Identifica se entramos na seção de avaliação do professor
        in_teacher_eval_section = False

        for idx, p in enumerate(self.doc.paragraphs):
            text = p.text.strip()
            if not text:
                continue

            text_lower = text.lower()
            if "uso do professor" in text_lower or "avaliação do professor" in text_lower:
                in_teacher_eval_section = True

            # Se for a seção do professor ou linhas de assinatura, ignora
            if ignore_signatures and in_teacher_eval_section:
                continue

            # Verifica se o parágrafo atual ou o seguinte é linha de assinatura
            is_sig = any(term in text_lower for term in ["assinatura", "carimbo"])
            if not is_sig and idx + 1 < len(self.doc.paragraphs):
                next_text = self.doc.paragraphs[idx + 1].text.lower()
                if any(term in next_text for term in ["assinatura", "carimbo"]):
                    is_sig = True

            if ignore_signatures and is_sig:
                continue

            # A. Verifica linhas de preenchimento (ex.: _____ ou ________ %)
            if check_underscores and "____" in text:
                report.add_issue(ValidationIssue(
                    category="placeholder",
                    location=f"Parágrafo #{idx}",
                    message="Linha pontilhada de preenchimento não respondida",
                    severity="error",
                    snippet=text[:80]
                ))

            # B. Verifica tags de template {{chave}} pendentes
            if check_tags:
                unfilled_tags = re.findall(r"\{\{([^}]+)\}\}", text)
                for tag in unfilled_tags:
                    report.add_issue(ValidationIssue(
                        category="tag",
                        location=f"Parágrafo #{idx}",
                        message=f"Tag de template pendente: '{{{{{tag}}}}}'",
                        severity="error",
                        snippet=text[:80]
                    ))

        # 2. Auditoria de Tabelas
        for t_idx, table in enumerate(self.doc.tables):
            if t_idx in ignore_tables:
                continue

            # Se for a última tabela de avaliação docente, pula
            if t_idx == len(self.doc.tables) - 1 and any("nota" in c.text.lower() for r in table.rows for c in r.cells):
                continue

            for r_idx, row in enumerate(table.rows):
                # Pula linha 0 se for cabeçalho
                if r_idx == 0 and len(table.rows) > 1:
                    continue

                for c_idx, cell in enumerate(row.cells):
                    cell_text = cell.text.strip()

                    # Verifica se a célula tem linhas de underscore
                    if check_underscores and "____" in cell_text:
                        report.add_issue(ValidationIssue(
                            category="placeholder",
                            location=f"Tabela #{t_idx} (Linha {r_idx}, Coluna {c_idx})",
                            message="Placeholder de preenchimento na célula",
                            severity="error",
                            snippet=cell_text[:50]
                        ))

                    # Verifica tags
                    if check_tags:
                        for tag in re.findall(r"\{\{([^}]+)\}\}", cell_text):
                            report.add_issue(ValidationIssue(
                                category="tag",
                                location=f"Tabela #{t_idx} (Linha {r_idx}, Coluna {c_idx})",
                                message=f"Tag pendente: '{{{{{tag}}}}}'",
                                severity="error",
                                snippet=cell_text[:50]
                            ))

                    # Verifica células vazias (quando ativado)
                    if check_empty_cells and not cell_text:
                        report.add_issue(ValidationIssue(
                            category="empty_cell",
                            location=f"Tabela #{t_idx} (Linha {r_idx}, Coluna {c_idx})",
                            message="Célula de dados está completamente vazia",
                            severity="warning",
                            snippet=""
                        ))

        return report
